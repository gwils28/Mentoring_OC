"""Dictionnaire des données de la base DATAImmo (livrable 1) : une entrée par colonne de la base."""

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from dataimmo.io import PROJECT_DIR

DICTIONNAIRE_PATH = PROJECT_DIR / "livrables" / "dictionnaire_donnees.xlsx"

VF = "Valeurs-foncières.xlsx"
REF = "fr-esr-referentiel-geographique.xlsx"
POP = "donnees_communes.xlsx"


@dataclass(frozen=True)
class Entree:
    table: str
    code: str
    signification: str
    type: str  # Texte, Entier, Décimal, Date (stockée en texte ISO dans SQLite)
    longueur: str
    nature: str  # Élémentaire, Calculée, Concaténée
    regle_gestion: str
    regle_calcul: str
    source: str


ENTREES = [
    # ---------------------------------------------------------------- region
    Entree("region", "id_region", "Code INSEE de la région", "Texte", "2", "Élémentaire",
           "Clé primaire, unique, non nulle ; texte pour garder le zéro initial (« 02 »)", "", f"{REF} : reg_code"),
    Entree("region", "nom_region", "Nom de la région", "Texte", "50", "Élémentaire",
           "Non nul ; écrit comme dans la source (« Ile-de-France » sans accent)", "", f"{REF} : reg_nom"),
    Entree("region", "nom_regroupement", "Grand ensemble : Province, Ile-de-France ou DROM-COM", "Texte", "20",
           "Élémentaire", "Non nul ; 3 valeurs possibles", "", f"{REF} : regrgp_nom"),
    # ---------------------------------------------------------------- commune
    Entree("commune", "id_codedep_codecommune", "Identifiant unique de la commune", "Texte", "6", "Concaténée",
           "Clé primaire, unique, non nulle", "code_departement + code_commune (« 34 » + « 172 » = « 34172 »)",
           f"{VF} : Code departement, Code commune"),
    Entree("commune", "id_region", "Région de la commune", "Texte", "2", "Élémentaire",
           "Clé étrangère vers region.id_region ; non nulle", "", f"{REF} : reg_code"),
    Entree("commune", "code_departement", "Code du département", "Texte", "3", "Élémentaire",
           "Non nul ; texte : « 2A », « 2B » (Corse), « 971 » (Guadeloupe), « 01 » garde son zéro", "",
           f"{VF} : Code departement"),
    Entree("commune", "code_commune", "Code de la commune dans son département", "Texte", "3", "Élémentaire",
           "Non nul ; toujours 3 caractères (« 004 »)", "", f"{VF} : Code commune"),
    Entree("commune", "nom_commune", "Nom officiel de la commune (la ville)", "Texte", "50", "Élémentaire",
           "Non nul ; encodage réparé (« Schœlcher »)", "", f"{REF} : com_nom"),
    Entree("commune", "population", "Population municipale de la commune", "Entier", "NC", "Élémentaire",
           "Non nulle, positive ou nulle", "", f"{POP} : PMUN"),
    # ---------------------------------------------------------------- bien
    Entree("bien", "id_bien", "Identifiant du bien dans la base", "Entier", "NC", "Calculée",
           "Clé primaire, unique, non nulle", "Numéro attribué automatiquement (1, 2, 3…)", "créé"),
    Entree("bien", "id_codedep_codecommune", "Commune où se trouve le bien", "Texte", "6", "Concaténée",
           "Clé étrangère vers commune ; non nulle", "code_departement + code_commune",
           f"{VF} : Code departement, Code commune"),
    Entree("bien", "no_voie", "Numéro dans la voie (adresse)", "Entier", "NC", "Élémentaire",
           "Peut être vide", "", f"{VF} : No voie"),
    Entree("bien", "btq", "Indice de répétition (B = bis, T = ter, Q = quater…)", "Texte", "1", "Élémentaire",
           "Peut être vide (94 % des biens)", "", f"{VF} : B/T/Q"),
    Entree("bien", "type_voie", "Type de voie (RUE, AV, BD, CHE…)", "Texte", "4", "Élémentaire",
           "Peut être vide ; abréviation de 4 caractères au plus", "", f"{VF} : Type de voie"),
    Entree("bien", "voie", "Nom de la voie (adresse)", "Texte", "50", "Élémentaire",
           "Non nul", "", f"{VF} : Voie"),
    Entree("bien", "code_postal", "Code postal de l'adresse", "Texte", "5", "Élémentaire",
           "5 caractères (zéro initial rétabli : « 01170 ») ; décrit l'adresse, pas la commune "
           "(Nice en a 4)", "", f"{VF} : Code postal"),
    Entree("bien", "type_local", "Type de bien", "Texte", "11", "Élémentaire",
           "Non nul ; « Appartement » ou « Maison »", "", f"{VF} : Type local"),
    Entree("bien", "total_pieces", "Nombre de pièces principales", "Entier", "NC", "Élémentaire",
           "Positif ou nul (0 existe dans la source)", "", f"{VF} : Nombre pieces principales"),
    Entree("bien", "surface_carrez", "Surface Carrez du 1er lot, en m²", "Décimal", "NC", "Élémentaire",
           "Strictement positive ; ne couvre que le 1er lot de la vente ; renseignée aussi pour les maisons "
           "dans ce fichier préparé", "", f"{VF} : Surface Carrez du 1er lot"),
    Entree("bien", "surface_reelle_bati", "Surface réelle bâtie, en m²", "Entier", "NC", "Élémentaire",
           "Positive ou nulle ; sert à contrôler la cohérence de la surface Carrez", "",
           f"{VF} : Surface reelle bati"),
    # ---------------------------------------------------------------- vente
    Entree("vente", "id_vente", "Identifiant de la vente dans la base", "Entier", "NC", "Calculée",
           "Clé primaire, unique, non nulle", "Numéro attribué automatiquement (1, 2, 3…)", "créé"),
    Entree("vente", "id_bien", "Bien vendu", "Entier", "NC", "Élémentaire",
           "Clé étrangère vers bien ; unique (1 vente par bien sur la période)", "", "créé"),
    Entree("vente", "date_vente", "Date de la vente (mutation)", "Date", "10", "Élémentaire",
           "Non nulle ; format AAAA-MM-JJ ; entre le 01/01/2020 et le 30/06/2020", "", f"{VF} : Date mutation"),
    Entree("vente", "valeur_fonciere", "Prix de vente en euros", "Décimal", "NC", "Élémentaire",
           "Strictement positive ; vide (NULL) pour 18 ventes non renseignées", "", f"{VF} : Valeur fonciere"),
    Entree("vente", "nombre_lots", "Nombre de lots compris dans la vente (cave, parking…)", "Entier", "NC",
           "Élémentaire", "Au moins 1", "", f"{VF} : Nombre de lots"),
]  # fmt: skip

LEXIQUE = [
    ("TABLE", "Table de la base où la donnée est rangée"),
    ("CODE", "Nom de la colonne dans la base"),
    ("SIGNIFICATION", "Ce que contient la colonne, en mots simples"),
    ("TYPE", "Texte, Entier, Décimal ou Date"),
    ("LONGUEUR", "Nombre maximal de caractères pour un texte ; NC (non concerné) pour un nombre"),
    (
        "NATURE",
        "Élémentaire (lue telle quelle), Calculée (produite par une règle) ou Concaténée (collage de plusieurs données)",
    ),
    ("RÈGLE DE GESTION", "Contrainte à respecter : unicité, non nulle, format, valeurs autorisées…"),
    ("RÈGLE DE CALCUL", "Comment la donnée est obtenue quand elle n'est pas élémentaire"),
    ("SOURCE", "Fichier et colonne d'origine (« créé » : n'existe pas dans les fichiers)"),
    (
        "Donnée retirée",
        "« Nom de l'acquéreur » : donnée personnelle sans utilité pour l'analyse, supprimée (RGPD)",
    ),
]


def ecrire_dictionnaire(path: Path = DICTIONNAIRE_PATH) -> Path:
    """Écrit le dictionnaire au format tableur : un onglet Dictionnaire, un onglet Lexique."""
    path.parent.mkdir(parents=True, exist_ok=True)
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Dictionnaire"
    entetes = ["TABLE", "CODE", "SIGNIFICATION", "TYPE", "LONGUEUR", "NATURE", "RÈGLE DE GESTION",
               "RÈGLE DE CALCUL", "SOURCE"]  # fmt: skip
    feuille.append(entetes)
    for e in ENTREES:
        feuille.append([e.table, e.code, e.signification, e.type, e.longueur, e.nature, e.regle_gestion,
                        e.regle_calcul, e.source])  # fmt: skip
    for cellule in feuille[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="123049")
    for colonne, largeur in zip("ABCDEFGHI", [10, 24, 40, 10, 10, 13, 55, 40, 38], strict=True):
        feuille.column_dimensions[colonne].width = largeur
    for ligne in feuille.iter_rows(min_row=2):
        for cellule in ligne:
            cellule.alignment = Alignment(wrap_text=True, vertical="top")
    feuille.freeze_panes = "C2"

    lexique = classeur.create_sheet("Lexique")
    lexique.append(["Terme", "Signification"])
    for ligne in LEXIQUE:
        lexique.append(list(ligne))
    lexique.column_dimensions["A"].width = 20
    lexique.column_dimensions["B"].width = 100
    for cellule in lexique[1]:
        cellule.font = Font(bold=True)
    classeur.save(path)
    return path
