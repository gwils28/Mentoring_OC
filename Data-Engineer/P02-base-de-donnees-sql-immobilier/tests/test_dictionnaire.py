"""Le dictionnaire doit décrire exactement la base : mêmes colonnes, types cohérents, longueurs tenues."""

import sqlite3
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from dataimmo.dictionnaire import ENTREES, ecrire_dictionnaire

PROJET = Path(__file__).resolve().parents[1]
TYPES_SQL = {"Texte": "TEXT", "Entier": "INTEGER", "Décimal": "REAL", "Date": "TEXT"}


def colonnes_du_schema() -> dict[tuple[str, str], str]:
    connexion = sqlite3.connect(":memory:")
    connexion.executescript((PROJET / "sql" / "schema.sql").read_text(encoding="utf-8"))
    colonnes = {}
    for (table,) in connexion.execute("SELECT name FROM sqlite_master WHERE type = 'table'"):
        for _, nom, type_sql, *_ in connexion.execute(f"PRAGMA table_info({table})"):
            colonnes[(table, nom)] = type_sql
    return colonnes


def test_le_dictionnaire_couvre_exactement_le_schema():
    decrites = {(e.table, e.code) for e in ENTREES}
    assert decrites == set(colonnes_du_schema())


def test_types_coherents_avec_le_schema():
    schema = colonnes_du_schema()
    for e in ENTREES:
        assert TYPES_SQL[e.type] == schema[(e.table, e.code)], f"{e.table}.{e.code}"


def test_variables_exigees_presentes():
    codes = {e.code for e in ENTREES}
    exigees = {
        "valeur_fonciere", "date_vente", "surface_carrez", "no_voie", "voie", "code_departement",
        "code_commune", "nom_commune", "total_pieces", "id_region", "nom_region",
        "id_codedep_codecommune", "type_local", "population",
    }  # fmt: skip
    assert exigees <= codes


def test_pas_de_donnee_personnelle():
    texte = " ".join(f"{e.code} {e.signification} {e.source}" for e in ENTREES).lower()
    assert "acqu" not in texte


def test_sept_colonnes_remplies_et_nature_valide():
    for e in ENTREES:
        assert e.signification and e.type and e.longueur and e.regle_gestion
        assert e.nature in {"Élémentaire", "Calculée", "Concaténée"}
        if e.nature != "Élémentaire":
            assert e.regle_calcul, (
                f"{e.code} : une donnée calculée ou concaténée doit avoir sa règle de calcul"
            )


@pytest.mark.skipif(not (PROJET / "data" / "processed" / "commune.csv").exists(), reason="CSV non construits")
def test_longueurs_tenues_par_les_donnees():
    for e in ENTREES:
        if e.type != "Texte" or not str(e.longueur).isdigit():
            continue
        valeurs = pd.read_csv(PROJET / "data" / "processed" / f"{e.table}.csv", dtype=str)[e.code].dropna()
        assert valeurs.str.len().max() <= int(e.longueur), f"{e.table}.{e.code}"


def test_fichier_excel(tmp_path):
    chemin = ecrire_dictionnaire(tmp_path / "dico.xlsx")
    classeur = openpyxl.load_workbook(chemin)
    assert classeur.sheetnames == ["Dictionnaire", "Lexique"]
    feuille = classeur["Dictionnaire"]
    entetes = [c.value for c in feuille[1]]
    assert entetes[:8] == [
        "TABLE", "CODE", "SIGNIFICATION", "TYPE", "LONGUEUR", "NATURE", "RÈGLE DE GESTION", "RÈGLE DE CALCUL",
    ]  # fmt: skip
    assert feuille.max_row == len(ENTREES) + 1
